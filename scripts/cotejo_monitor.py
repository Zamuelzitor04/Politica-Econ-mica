"""Cotejo con la base enviada por el monitor y tabla de unidades (respuesta a sus dos advertencias).

Advertencias del monitor (22-oct-2026):
  1. Un aumento de la cartera en pesos no implica una expansión igual de su poder adquisitivo → deflactar.
  2. Las series tienen unidades y años de referencia distintos → armonizar antes de comparar.

Lee data/raw/base_de_datos_taller_2_monitor.xlsx (sin modificarlo) y escribe dos hojas en el Excel:
  - Cotejo_monitor: nuestras series vs. las del monitor, con diferencias calculadas por fórmula.
  - Unidades: unidad y referencia de cada serie, conversión aplicada y un ejemplo nominal vs. real.
"""
from __future__ import annotations

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from construir_base import RAW

ARCH_MONITOR = RAW / "base_de_datos_taller_2_monitor.xlsx"


def _hoja(wb, nombre) -> pd.DataFrame:
    filas = [f for f in wb[nombre].iter_rows(values_only=True) if any(c is not None for c in f)]
    return pd.DataFrame(filas[1:], columns=filas[0])


def leer_monitor() -> dict:
    wb = openpyxl.load_workbook(ARCH_MONITOR, data_only=True)
    tpm = _hoja(wb, "tasa política")
    tpm["fecha"] = pd.to_datetime(tpm["fecha"])
    mensual = {}
    for hoja, clave in (("tasa consumo", "tasa_consumo"), ("cartera vencida", "vencida_millones"),
                        ("ipc", "ipc"), ("inflacion", "inflacion_anual")):
        h = _hoja(wb, hoja)
        idx = [pd.Period(year=int(a), month=int(m), freq="M") for a, m in zip(h.iloc[:, 1], h.iloc[:, 2])]
        mensual[clave] = pd.Series(h.iloc[:, 3].astype(float).values, index=idx)
    cons = _hoja(wb, "consumo hogares").set_index("Período")
    return {
        "tpm_diaria": tpm.set_index("fecha")["tasa de intervencion"].astype(float),
        "mensual": pd.DataFrame(mensual),
        "consumo": cons.iloc[:, 2].astype(float),
        "estado": cons["Estado de la cifra"],
        "hojas": wb.sheetnames,
    }


def resumen(mon: dict, diaria: pd.DataFrame, m: pd.DataFrame, q: pd.DataFrame) -> dict:
    """Diferencias máximas absolutas (para imprimir al construir la base)."""
    o = diaria[diaria["fecha"] <= "2026-06-30"].set_index("fecha")["tpm"]
    mm = mon["mensual"]
    return {
        "TPM diaria": (mon["tpm_diaria"] - o).abs().max(),
        "Tasa de consumo": (mm["tasa_consumo"] - m["tasa_consumo"]).abs().max(),
        "Cartera vencida (monitor/1000)": (mm["vencida_millones"] / 1000 - m["cartera_vencida"]).abs().max(),
        "IPC": (mm["ipc"] - m["ipc"]).abs().max(),
        "Inflación anual": (mm["inflacion_anual"] - m["inflacion_anual"]).abs().max(),
        "Consumo hogares": (mon["consumo"] - q["consumo_hogares"]).abs().max(),
    }


def hoja_cotejo(wb, mon: dict, fila_mes: dict, f0_base: int, q: pd.DataFrame):
    from exportar import AZUL_DATO, CENTRO, GRIS, NEGRITA, NEGRO, RELLENO_ENC, TITULO, VERDE, _fmt

    ws = wb.create_sheet("Cotejo_monitor", 5)
    _fmt(ws, "A1", "Cotejo con la base enviada por el monitor (base_de_datos_taller_2.xlsx, 22-oct-2026)", TITULO)
    _fmt(ws, "A2", "Verde = nuestra base (vínculo); azul = dato del monitor; negro = diferencia. Diferencia 0 = "
                   "misma cifra. La base del monitor NO trae cartera bruta: se toma de la SFC (hoja C5), misma "
                   "fuente y cobertura que su cartera vencida.", GRIS)
    enc = ["Mes", "TPM nuestra (prom. diario, %)", "TPM monitor (prom. diario, %)", "Dif.",
           "Tasa consumo nuestra (%)", "Tasa consumo monitor (%)", "Dif.",
           "IPC nuestro", "IPC monitor", "Dif.", "Inflación nuestra (%)", "Inflación monitor (%)", "Dif.",
           "Vencida nuestra (miles de millones $)", "Vencida monitor (millones $)",
           "Vencida monitor ÷ 1.000 (miles de millones $)", "Dif."]
    for j, t in enumerate(enc, start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.row_dimensions[4].height = 72

    # Serie diaria del monitor en columnas X:Y (insumo del promedio mensual)
    _fmt(ws, "X4", "Fecha (monitor)", NEGRITA)
    _fmt(ws, "Y4", "TPM diaria monitor (%)", NEGRITA)
    td = mon["tpm_diaria"]
    for i, (fch, v) in enumerate(td.items(), start=5):
        _fmt(ws, f"X{i}", fch.to_pydatetime(), AZUL_DATO, "dd/mm/yyyy")
        _fmt(ws, f"Y{i}", float(v), AZUL_DATO, "0.00")
    ult_d = 4 + len(td)
    rf, rv = f"$X$5:$X${ult_d}", f"$Y$5:$Y${ult_d}"

    mm = mon["mensual"]
    r = 5
    for mes, fm in fila_mes.items():
        _fmt(ws, f"A{r}", f"=Mensual!A{fm}", VERDE, "mmm-yyyy")
        _fmt(ws, f"B{r}", f"=Mensual!C{fm}", VERDE, "0.000")
        _fmt(ws, f"C{r}", f'=AVERAGEIFS({rv},{rf},">="&A{r},{rf},"<="&EOMONTH(A{r},0))', NEGRO, "0.000")
        _fmt(ws, f"E{r}", f"=Mensual!E{fm}", VERDE, "0.00")
        _fmt(ws, f"F{r}", float(mm.loc[mes, "tasa_consumo"]), AZUL_DATO, "0.00")
        _fmt(ws, f"H{r}", f"=Mensual!F{fm}", VERDE, "0.00")
        _fmt(ws, f"I{r}", float(mm.loc[mes, "ipc"]), AZUL_DATO, "0.00")
        _fmt(ws, f"K{r}", f"=Mensual!G{fm}", VERDE, "0.00")
        _fmt(ws, f"L{r}", float(mm.loc[mes, "inflacion_anual"]), AZUL_DATO, "0.00")
        _fmt(ws, f"N{r}", f"=Mensual!I{fm}", VERDE, "#,##0.000")
        _fmt(ws, f"O{r}", float(mm.loc[mes, "vencida_millones"]), AZUL_DATO, "#,##0.000")
        _fmt(ws, f"P{r}", f"=O{r}/1000", NEGRO, "#,##0.000")
        for dif, a, b in (("D", "B", "C"), ("G", "E", "F"), ("J", "H", "I"), ("M", "K", "L"), ("Q", "N", "P")):
            _fmt(ws, f"{dif}{r}", f"={a}{r}-{b}{r}", NEGRO, "0.000000")
        r += 1
    ult = r - 1
    _fmt(ws, f"A{r + 1}", "Máx. |dif.|", NEGRITA)
    for col in "DGJMQ":
        _fmt(ws, f"{col}{r + 1}", f"=MAX(MAX({col}5:{col}{ult}),-MIN({col}5:{col}{ult}))", NEGRITA, "0.000000")

    # Consumo de los hogares (trimestral)
    r0 = r + 4
    _fmt(ws, f"A{r0 - 1}", "Consumo final de los hogares (miles de millones $, volumen referencia 2015)", NEGRITA)
    for j, t in enumerate(["Trimestre", "Nuestro", "Monitor", "Dif.", "Estado de la cifra (monitor)"], start=1):
        c = ws.cell(r0, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
    for k, trim in enumerate(q.index):
        rr = r0 + 1 + k
        _fmt(ws, f"A{rr}", trim, NEGRO)
        _fmt(ws, f"B{rr}", f"=Base_trimestral!N{f0_base + k}", VERDE, "#,##0.000")
        _fmt(ws, f"C{rr}", float(mon["consumo"].loc[trim]), AZUL_DATO, "#,##0.000")
        _fmt(ws, f"D{rr}", f"=B{rr}-C{rr}", NEGRO, "0.000000")
        _fmt(ws, f"E{rr}", str(mon["estado"].loc[trim]), AZUL_DATO)
    ws.freeze_panes = "B5"


def hoja_unidades(wb, f0_base: int, q: pd.DataFrame):
    from exportar import AJUSTE, CENTRO, GRIS, NEGRITA, NEGRO, RELLENO_ENC, TITULO, VERDE, _fmt

    ws = wb.create_sheet("Unidades", 2)
    _fmt(ws, "A1", "Unidades, años de referencia y cómo se hacen comparables las series", TITULO)
    _fmt(ws, "A2", "Responde a las dos advertencias del monitor: (1) pesos corrientes ≠ poder adquisitivo; "
                   "(2) unidades y años de referencia distintos.", GRIS)
    enc = ["Serie", "Unidad en la base del monitor", "Unidad en la fuente original", "Unidad en nuestra base",
           "Año / período de referencia", "Conversión aplicada", "Cómo se compara con las demás"]
    anchos = [20, 26, 28, 28, 24, 34, 44]
    filas = [
        ("TPM", "% (nominal, efectivo anual)", "% (BanRep, dato diario)", "% e.a., promedio trimestral",
         "No aplica (tasa)", "Promedio diario por mes → promedio de 3 meses",
         "Misma unidad que la tasa de consumo: se grafican juntas (Fig. 1). En términos reales: Fig. E1."),
        ("Tasa de crédito de consumo", "% (nominal, efectivo anual)", "% e.a., promedio mensual (BanRep)",
         "% e.a., promedio trimestral", "No aplica (tasa)", "Promedio de 3 meses",
         "Comparable con la TPM. Excluye tarjetas de crédito; la cartera SFC las incluye."),
        ("Cartera vencida (V)", "Millones de COP corrientes", "Miles de millones de COP corrientes (SFC C5)",
         "Miles de millones de COP; reales en pesos de jun-2026",
         "Pesos corrientes de cada mes", "Monitor ÷ 1.000 = nuestra cifra. Real = V × IPC(jun-2026) / IPC(mes)",
         "Nunca comparar pesos de meses distintos sin deflactar. Índice real 2021-I = 100 (Fig. 2)."),
        ("Cartera bruta (B)", "No viene en la base del monitor", "Miles de millones de COP corrientes (SFC C5)",
         "Miles de millones de COP; reales en pesos de jun-2026",
         "Pesos corrientes de cada mes", "Real = B × IPC(jun-2026) / IPC(mes)",
         "Igual que V. La morosidad (100 × V / B) no necesita deflactar: el IPC se cancela."),
        ("Consumo final de los hogares", "Miles de millones de COP, referencia 2015",
         "Miles de millones de COP, series encadenadas de volumen, año de referencia 2015 (DANE)",
         "Igual; además índice 2021-I = 100", "Volumen con referencia 2015 (precios de 2015 encadenados)",
         "Ninguna (ya es real). No se deflacta otra vez.",
         "Sus 'pesos de 2015' no son los pesos de jun-2026 de la cartera: no se comparan niveles en pesos, "
         "solo índices 2021-I = 100 o tasas de crecimiento."),
        ("IPC", "Índice, dic-2018 = 100", "Índice, dic-2018 = 100 (DANE)", "Igual",
         "Base diciembre de 2018", "Se usa como cociente IPC(jun-2026) / IPC(mes)",
         "La base 2018 desaparece en el cociente: lo que importa es cuánto subieron los precios entre dos meses."),
        ("Inflación anual", "%", "% variación 12 meses (DANE)", "Igual (dato del último mes del trimestre)",
         "No aplica", "Ninguna", "Contexto; también sirve para tasas reales (Fig. E1)."),
    ]
    for j, (t, w) in enumerate(zip(enc, anchos), start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, fila in enumerate(filas, start=5):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(i, j, v)
            c.font, c.alignment = NEGRO, AJUSTE
        ws.row_dimensions[i].height = 75

    # Ejemplo: nominal vs. real, 2021-I → 2026-II (fórmulas sobre Base_trimestral)
    r = 5 + len(filas) + 2
    fi = f0_base + list(q.index).index("2021-I")
    ff = f0_base + list(q.index).index("2026-II")
    _fmt(ws, f"A{r}", "Ejemplo (advertencia 1): cambio de 2021-I a 2026-II", NEGRITA)
    for j, t in enumerate(["Concepto", "2021-I", "2026-II", "Cambio (%)"], start=1):
        c = ws.cell(r + 1, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
    ejemplo = [("Cartera bruta en pesos corrientes (miles de millones)", "G"),
               ("IPC (nivel de precios)", "F"),
               ("Cartera bruta real (miles de millones de pesos de jun-2026)", "I"),
               ("Cartera vencida en pesos corrientes (miles de millones)", "H"),
               ("Cartera vencida real (miles de millones de pesos de jun-2026)", "J")]
    for k, (txt, col) in enumerate(ejemplo):
        rr = r + 2 + k
        _fmt(ws, f"A{rr}", txt, NEGRO)
        _fmt(ws, f"B{rr}", f"=Base_trimestral!{col}{fi}", VERDE, "#,##0.00")
        _fmt(ws, f"C{rr}", f"=Base_trimestral!{col}{ff}", VERDE, "#,##0.00")
        _fmt(ws, f"D{rr}", f"=100*(C{rr}/B{rr}-1)", NEGRO, "0.0")
    _fmt(ws, f"A{r + 8}", "Lectura: en pesos la cartera bruta creció cerca de 39 %, pero los precios subieron cerca de "
                         "49 %; su poder adquisitivo cayó cerca de 7 %. La vencida subió cerca de 14 % en pesos y "
                         "cayó cerca de 24 % en términos reales.", GRIS)
