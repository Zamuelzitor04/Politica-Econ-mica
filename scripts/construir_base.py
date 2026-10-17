"""Construye la base trimestral del Taller 2 (crédito, consumo y morosidad).

Lee únicamente los archivos originales de data/raw/ (sin modificarlos) y produce:
  - data/processed/insumos_mensuales.csv
  - data/processed/base_trimestral.csv
  - output/figuras/figura{1..4}.png
  - output/Taller2_base_trimestral.xlsx  (base con fórmulas, ficha de fuentes, controles y figuras)

Uso:  python scripts/construir_base.py
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[1]
RAW = RAIZ / "data" / "raw"
PROC = RAIZ / "data" / "processed"
OUT = RAIZ / "output"
FIG = OUT / "figuras"

ARCH_TASAS = RAW / "Datos_Adicionales_BANREP.xlsx"
ARCH_IPC = RAW / "Datos_BANREP_2.xlsx"
ARCH_SFC = RAW / "Cartera_hasta_agosto_Superfinanciera.xlsx"
ARCH_DANE = RAW / "anex-GastoConstantes-IItrim2026.xlsx"

INICIO = pd.Period("2021-01", "M")
FIN = pd.Period("2026-06", "M")
MES_BASE_IPC = pd.Period("2026-06", "M")  # saldos reales en pesos de junio de 2026
TRIM_BASE = "2021-I"                      # base de los índices (= 100)
ROMANOS = {1: "I", 2: "II", 3: "III", 4: "IV"}


def etiqueta_trim(p: pd.Period) -> str:
    return f"{p.year}-{ROMANOS[p.quarter]}"


def num_banrep(x):
    """Las celdas del BanRep vienen como texto '1.234,56'; '-' significa sin dato."""
    if x is None or str(x).strip() in ("", "-"):
        return None
    return float(str(x).replace(".", "").replace(",", "."))


# --------------------------------------------------------------------------- lectura
def leer_banrep(path: Path) -> pd.DataFrame:
    filas = list(openpyxl.load_workbook(path, data_only=True)["Datos"].iter_rows(values_only=True))
    encabezado = filas[0]
    datos = [f for f in filas[2:] if f[0] and len(str(f[0])) == 10 and str(f[0])[2] == "/"]
    df = pd.DataFrame(datos, columns=encabezado)
    df["Fecha"] = pd.to_datetime(df["Fecha"], format="%d/%m/%Y")
    for c in df.columns[1:]:
        df[c] = df[c].map(num_banrep)
    return df


def leer_tasas():
    df = leer_banrep(ARCH_TASAS)
    diaria = df[["Fecha", "Tasa de política monetaria(Dato diario)"]].dropna()
    diaria.columns = ["fecha", "tpm"]
    dato_mes = df[["Fecha", "Tasa de política monetaria(Dato mes)"]].dropna()
    dato_mes = pd.Series(dato_mes.iloc[:, 1].values, index=dato_mes["Fecha"].dt.to_period("M"))
    consumo = df[["Fecha", "Crédito de consumo, Tasa de interés(Promedio mensual)"]].dropna()
    consumo = pd.Series(consumo.iloc[:, 1].values, index=consumo["Fecha"].dt.to_period("M"))
    return diaria, dato_mes, consumo


def leer_ipc():
    df = leer_banrep(ARCH_IPC)
    df["mes"] = df["Fecha"].dt.to_period("M")
    df = df.set_index("mes")
    ipc = df["Índice de Precios al Consumidor (IPC)(Dato fin de mes)"].dropna()
    infl = df["Índice de Precios al Consumidor (IPC)(Variación porcentual  anual)"].dropna()
    return ipc, infl


def leer_sfc() -> pd.DataFrame:
    filas = list(openpyxl.load_workbook(ARCH_SFC, data_only=True)["C5"].iter_rows(values_only=True))
    enc = list(filas[3])
    ib, iv, ic = enc.index("CONSUMO - BRUTA"), enc.index("CONSUMO - VENCIDA"), enc.index(
        "CONSUMO - INDICADOR CALIDAD TRADICIONAL")
    reg = []
    for f in filas[4:]:
        try:
            mes = pd.Period(dt.datetime.strptime(str(f[0]), "%B-%Y"), "M")
        except ValueError:
            continue
        reg.append((mes, float(f[ib]), float(f[iv]), float(f[ic])))
    return pd.DataFrame(reg, columns=["mes", "bruta", "vencida", "ind_sfc"]).set_index("mes")


def leer_dane() -> pd.DataFrame:
    filas = list(openpyxl.load_workbook(ARCH_DANE, data_only=True)["Cuadro 2"].iter_rows(values_only=True))
    concepto = "Gasto de consumo final individual de los hogares"
    idx = [i for i, f in enumerate(filas) if f[1] and str(f[1]).startswith(concepto)]
    i_nivel, i_var = idx[0], idx[1]          # 1ª tabla: niveles; 2ª tabla: crecimiento trimestral
    i_anio = max(i for i in range(i_nivel) if filas[i][0] == "Clasificación Cuentas Nacionales")
    anios, trims = filas[i_anio], filas[i_anio + 1]
    reg, anio = [], None
    for j in range(2, len(trims)):
        if anios[j]:
            anio = int(str(anios[j])[:4])
        if trims[j] is None:
            continue
        q = {"I": 1, "II": 2, "III": 3, "IV": 4}[trims[j]]
        reg.append((pd.Period(year=anio, quarter=q, freq="Q"), float(filas[i_nivel][j]),
                    None if filas[i_var][j] is None else float(filas[i_var][j])))
    return pd.DataFrame(reg, columns=["trim", "consumo", "var_dane"]).set_index("trim")


# --------------------------------------------------------------------------- cálculo
def construir():
    diaria, dato_mes, tasa_cons = leer_tasas()
    ipc, infl = leer_ipc()
    sfc = leer_sfc()
    dane = leer_dane()

    meses = pd.period_range(INICIO, FIN, freq="M")
    tpm_mes = diaria.assign(mes=diaria["fecha"].dt.to_period("M")).groupby("mes")["tpm"].mean()
    m = pd.DataFrame(index=meses)
    m["tpm_prom_diario"] = tpm_mes
    m["tpm_dato_mes"] = dato_mes
    m["tasa_consumo"] = tasa_cons
    m["ipc"] = ipc
    m["inflacion_anual"] = infl
    m["cartera_bruta"] = sfc["bruta"]
    m["cartera_vencida"] = sfc["vencida"]
    m["ind_calidad_sfc"] = sfc["ind_sfc"]
    m["trimestre"] = [etiqueta_trim(p.asfreq("Q")) for p in m.index]

    ipc_base = ipc[MES_BASE_IPC]
    trims = pd.period_range("2020Q4", "2026Q2", freq="Q")
    q = pd.DataFrame(index=[etiqueta_trim(t) for t in trims])
    for t in trims:
        e = etiqueta_trim(t)
        if t < pd.Period("2021Q1", "Q"):
            continue  # 2020-IV: solo consumo
        mm = m[m["trimestre"] == e]
        ult = mm.index.max()
        q.loc[e, "tpm"] = mm["tpm_prom_diario"].mean()
        q.loc[e, "tasa_consumo"] = mm["tasa_consumo"].mean()
        q.loc[e, "ipc"] = m.loc[ult, "ipc"]
        q.loc[e, "inflacion_anual"] = m.loc[ult, "inflacion_anual"]
        q.loc[e, "bruta_nom"] = m.loc[ult, "cartera_bruta"]
        q.loc[e, "vencida_nom"] = m.loc[ult, "cartera_vencida"]
    q["diferencial"] = q["tasa_consumo"] - q["tpm"]
    q["bruta_real"] = q["bruta_nom"] * ipc_base / q["ipc"]
    q["vencida_real"] = q["vencida_nom"] * ipc_base / q["ipc"]
    q["morosidad"] = 100 * q["vencida_nom"] / q["bruta_nom"]
    q["idx_bruta_real"] = 100 * q["bruta_real"] / q.loc[TRIM_BASE, "bruta_real"]
    q["idx_vencida_real"] = 100 * q["vencida_real"] / q.loc[TRIM_BASE, "vencida_real"]
    q["consumo_hogares"] = [dane.loc[t, "consumo"] if t in dane.index else None for t in trims]
    q["idx_consumo"] = 100 * q["consumo_hogares"] / q.loc[TRIM_BASE, "consumo_hogares"]
    q["var_trim_consumo"] = 100 * (q["consumo_hogares"] / q["consumo_hogares"].shift(1) - 1)
    q["var_dane"] = [dane.loc[t, "var_dane"] if t in dane.index else None for t in trims]
    return diaria, m, q, ipc_base


if __name__ == "__main__":
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from exportar import exportar

    diaria, mensual, trimestral, ipc_base = construir()
    exportar(diaria, mensual, trimestral, ipc_base)
