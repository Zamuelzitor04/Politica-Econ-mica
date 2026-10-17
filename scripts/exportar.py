"""Exporta la base del Taller 2: CSV, figuras PNG y libro Excel con fórmulas.

Lo llama scripts/construir_base.py. En el Excel los datos transcritos de las fuentes van en azul
y todo lo calculado es fórmula, para que cada cifra se pueda rastrear en la sustentación.
"""
from __future__ import annotations

import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from construir_base import FIG, MES_BASE_IPC, OUT, PROC, TRIM_BASE, etiqueta_trim

# ----------------------------------------------------------------------------- estilo
AZUL_DATO = Font(name="Arial", size=10, color="0000FF")
NEGRO = Font(name="Arial", size=10)
VERDE = Font(name="Arial", size=10, color="008000")
NEGRITA = Font(name="Arial", size=10, bold=True)
TITULO = Font(name="Arial", size=13, bold=True)
GRIS = Font(name="Arial", size=9, italic=True, color="52514E")
RELLENO_ENC = PatternFill("solid", fgColor="E8EEF7")
RELLENO_PARAM = PatternFill("solid", fgColor="FFFF00")
BORDE = Border(bottom=Side(style="thin", color="999999"))
AJUSTE = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

SERIE1, SERIE2 = "#2a78d6", "#eb6834"  # paleta categórica validada (azul, naranja)
TINTA, TINTA2, REJILLA = "#0b0b0b", "#52514e", "#dddcd7"

FUENTE_TASAS = ("Fuente: Banco de la República, con información de la Superintendencia Financiera "
                "(tasa de consumo). Cálculos propios.")
FUENTE_SFC = "Fuente: Superintendencia Financiera de Colombia (cartera); DANE (IPC). Cálculos propios."
FUENTE_DANE = "Fuente: DANE, PIB trimestral por el enfoque del gasto (actualizado 18-ago-2026). Cálculos propios."


def _fmt(ws, celda, valor, fuente=NEGRO, formato=None):
    c = ws[celda]
    c.value = valor
    c.font = fuente
    if formato:
        c.number_format = formato
    return c


# ----------------------------------------------------------------------------- figuras
def _ejes(ax, etiquetas):
    ax.set_xticks(range(len(etiquetas)))
    ax.set_xticklabels(etiquetas, rotation=45, ha="right", fontsize=8, color=TINTA2)
    ax.tick_params(axis="y", labelsize=9, colors=TINTA2)
    ax.grid(axis="y", color=REJILLA, linewidth=0.8)
    ax.set_axisbelow(True)
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    for lado in ("left", "bottom"):
        ax.spines[lado].set_color("#b5b4ad")


def _coma(v, dec):
    """Formato numérico con coma decimal (convención en Colombia)."""
    return f"{v:.{dec}f}".replace(".", ",")


def _marco(fig, titulo, subtitulo, fuente, nota):
    fig.suptitle(titulo, x=0.01, y=0.985, ha="left", fontsize=13, fontweight="bold", color=TINTA)
    fig.text(0.01, 0.925, subtitulo, ha="left", fontsize=10, color=TINTA2)
    pie = fuente + "\n" + "\n".join(textwrap.wrap("Nota: " + nota, 160))
    fig.text(0.01, 0.012, pie, ha="left", va="bottom", fontsize=8, color=TINTA2, linespacing=1.5)


def _etiqueta_final(ax, x, y, texto, color):
    ax.annotate(texto, (x, y), xytext=(6, 0), textcoords="offset points", va="center",
                fontsize=9, color=TINTA, fontweight="bold")
    ax.plot([x], [y], "o", color=color, markersize=5, markeredgecolor="white", markeredgewidth=1.5)


def figuras(q: pd.DataFrame) -> list:
    FIG.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"
    d = q.loc[TRIM_BASE:]
    et = list(d.index)
    x = range(len(et))
    rutas = []

    # Figura 1 ---------------------------------------------------------------
    fig, ax = plt.subplots(figsize=(10, 5.8))
    fig.subplots_adjust(top=0.86, bottom=0.2, left=0.07, right=0.93)
    ax.plot(x, d["tpm"], color=SERIE1, lw=2, label="Tasa de intervención (TPM)")
    ax.plot(x, d["tasa_consumo"], color=SERIE2, lw=2, label="Tasa de crédito de consumo (sin tarjetas)")
    _etiqueta_final(ax, len(et) - 1, d["tpm"].iloc[-1], _coma(d['tpm'].iloc[-1], 2), SERIE1)
    _etiqueta_final(ax, len(et) - 1, d["tasa_consumo"].iloc[-1], _coma(d['tasa_consumo'].iloc[-1], 2), SERIE2)
    i_max = d["tasa_consumo"].values.argmax()
    ax.annotate(f"Máximo {_coma(d['tasa_consumo'].iloc[i_max], 2)} ({et[i_max]})", (i_max, d["tasa_consumo"].iloc[i_max]),
                xytext=(10, 6), textcoords="offset points", fontsize=8.5, color=TINTA2)
    _ejes(ax, et)
    ax.set_ylabel("% efectivo anual", color=TINTA2)
    ax.set_ylim(0, 33)
    ax.legend(loc="upper right", frameon=False, fontsize=9)
    _marco(fig, "Figura 1. Tasa de intervención y tasa del crédito de consumo",
           "Promedio trimestral de tasas mensuales, % efectivo anual, 2021-I a 2026-II", FUENTE_TASAS,
           "TPM: promedio de la serie diaria de cada mes y luego de los tres meses. Tasa de consumo: promedio "
           "de las tasas mensuales (créditos nuevos, ponderadas por monto; no incluye tarjetas de crédito).")
    rutas.append(FIG / "figura1_tasas.png")
    fig.savefig(rutas[-1], dpi=200)
    plt.close(fig)

    # Figura 2 ---------------------------------------------------------------
    fig, ax = plt.subplots(figsize=(10, 5.8))
    fig.subplots_adjust(top=0.86, bottom=0.2, left=0.07, right=0.93)
    ax.axhline(100, color="#8a8984", lw=1, ls="--")
    ax.plot(x, d["idx_bruta_real"], color=SERIE1, lw=2, label="Cartera bruta de consumo, real")
    ax.plot(x, d["idx_vencida_real"], color=SERIE2, lw=2, label="Cartera vencida de consumo, real")
    _etiqueta_final(ax, len(et) - 1, d["idx_bruta_real"].iloc[-1], _coma(d['idx_bruta_real'].iloc[-1], 1), SERIE1)
    _etiqueta_final(ax, len(et) - 1, d["idx_vencida_real"].iloc[-1], _coma(d['idx_vencida_real'].iloc[-1], 1), SERIE2)
    i_max = d["idx_vencida_real"].values.argmax()
    ax.annotate(f"{_coma(d['idx_vencida_real'].iloc[i_max], 1)} ({et[i_max]})", (i_max, d["idx_vencida_real"].iloc[i_max]),
                xytext=(0, 8), textcoords="offset points", ha="center", fontsize=8.5, color=TINTA2)
    _ejes(ax, et)
    ax.set_ylabel("Índice, 2021-I = 100", color=TINTA2)
    ax.legend(loc="upper left", frameon=False, fontsize=9)
    _marco(fig, "Figura 2. Cartera bruta y vencida de consumo en términos reales",
           "Índices, 2021-I = 100; saldos del último mes de cada trimestre, 2021-I a 2026-II", FUENTE_SFC,
           "Establecimientos de crédito sin IOE, incluye leasing. Saldos deflactados con el IPC del mismo mes "
           "(base jun-2026). Vencida: créditos con mora, según la SFC.")
    rutas.append(FIG / "figura2_cartera_real.png")
    fig.savefig(rutas[-1], dpi=200)
    plt.close(fig)

    # Figura 3 ---------------------------------------------------------------
    fig, ax = plt.subplots(figsize=(10, 5.8))
    fig.subplots_adjust(top=0.86, bottom=0.2, left=0.07, right=0.93)
    ax.plot(x, d["morosidad"], color=SERIE1, lw=2)
    for i in sorted({0, int(d["morosidad"].values.argmax()), int(d["morosidad"].values.argmin()), len(et) - 1}):
        ax.plot([i], [d["morosidad"].iloc[i]], "o", color=SERIE1, markersize=5, markeredgecolor="white",
                markeredgewidth=1.5)
        ax.annotate(f"{_coma(d['morosidad'].iloc[i], 2)} %", (i, d["morosidad"].iloc[i]), xytext=(0, 9),
                    textcoords="offset points", ha="center", fontsize=9, color=TINTA, fontweight="bold")
    _ejes(ax, et)
    ax.set_ylabel("% de la cartera bruta", color=TINTA2)
    ax.set_ylim(0, 10)
    _marco(fig, "Figura 3. Índice de morosidad de la cartera de consumo",
           "Cartera vencida / cartera bruta × 100, último mes de cada trimestre, 2021-I a 2026-II", FUENTE_SFC,
           "Establecimientos de crédito sin IOE, incluye leasing. Coincide con el indicador de calidad "
           "tradicional publicado por la SFC. Puede cambiar por el saldo vencido, la cartera bruta o ambos.")
    rutas.append(FIG / "figura3_morosidad.png")
    fig.savefig(rutas[-1], dpi=200)
    plt.close(fig)

    # Figura 4 ---------------------------------------------------------------
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(10, 8), sharex=True)
    fig.subplots_adjust(top=0.865, bottom=0.15, left=0.08, right=0.93, hspace=0.25)
    a1.axhline(100, color="#8a8984", lw=1, ls="--")
    a1.plot(x, d["idx_consumo"], color=SERIE1, lw=2)
    _etiqueta_final(a1, len(et) - 1, d["idx_consumo"].iloc[-1], _coma(d['idx_consumo'].iloc[-1], 1), SERIE1)
    a1.set_title("A. Índice del consumo real de los hogares (2021-I = 100)", loc="left", fontsize=10,
                 color=TINTA)
    a1.set_ylabel("Índice", color=TINTA2)
    colores = [SERIE1 if v >= 0 else SERIE2 for v in d["var_trim_consumo"]]
    a2.bar(x, d["var_trim_consumo"], color=colores, width=0.7)
    a2.axhline(0, color="#8a8984", lw=1)
    a2.set_ylim(d["var_trim_consumo"].min() - 0.8, d["var_trim_consumo"].max() + 0.8)
    for i, v in enumerate(d["var_trim_consumo"]):
        a2.annotate(_coma(v, 1), (i, v), xytext=(0, 3 if v >= 0 else -10), textcoords="offset points",
                    ha="center", fontsize=7.5, color=TINTA2)
    a2.set_title("B. Variación trimestral del consumo real (%); naranja = caída", loc="left", fontsize=10,
                 color=TINTA)
    a2.set_ylabel("% frente al trimestre anterior", color=TINTA2)
    _ejes(a1, et)
    _ejes(a2, et)
    _marco(fig, "Figura 4. Consumo real de los hogares",
           "Volumen (series encadenadas, referencia 2015), ajustado por estacionalidad y calendario, 2021-I a 2026-II",
           FUENTE_DANE,
           "Incluye ISFLSH. No se deflacta de nuevo. Variación = 100 × (C_t / C_t-1 − 1); 2021-I usa 2020-IV. "
           "2024 provisional; 2025-2026 preliminar.")
    rutas.append(FIG / "figura4_consumo.png")
    fig.savefig(rutas[-1], dpi=200)
    plt.close(fig)
    return rutas


# ----------------------------------------------------------------------------- Excel
FICHA = [
    ("Tasa de intervención (TPM)", "Banco de la República",
     "Tasa de política monetaria (Dato diario)",
     "Tasa mínima de las subastas de expansión a un día (OMA)", "% (efectivo anual)",
     "Diaria, 01/01/2021–30/06/2026", "Datos_Adicionales_BANREP.xlsx, hoja Datos",
     "https://www.banrep.gov.co/es/estadisticas  [pegar la URL exacta de la consulta]",
     "16-oct-2026 (pie del archivo)",
     "Promedio simple de los días de cada mes → promedio simple de los 3 meses del trimestre."),
    ("Tasa del crédito de consumo", "Banco de la República con información de la SFC",
     "Crédito de consumo, Tasa de interés (Promedio mensual)",
     "Créditos de consumo desembolsados; no incluye tarjetas de crédito", "% efectivo anual",
     "Mensual, ene-2021–jun-2026", "Datos_Adicionales_BANREP.xlsx, hoja Datos",
     "https://www.banrep.gov.co/es/estadisticas  [pegar la URL exacta de la consulta]",
     "16-oct-2026 (pie del archivo)",
     "Promedio simple de las 3 tasas mensuales del trimestre."),
    ("Cartera bruta de consumo (B)", "Superintendencia Financiera de Colombia",
     "Evolución de la cartera de crédito por tipo de crédito — CONSUMO - BRUTA (hoja C5)",
     "Establecimientos de crédito, sin Instituciones Oficiales Especiales; incluye leasing",
     "Miles de millones de pesos corrientes (el título de C5 dice 'millones'; C5 suma exactamente C4, "
     "rotulada en miles de millones)",
     "Mensual, ene-2016–ago-2026; se usan mar/jun/sep/dic 2021-2026",
     "Cartera_hasta_agosto_Superfinanciera.xlsx, hoja C5",
     "https://www.superfinanciera.gov.co  (Evolución de la cartera de créditos)  [pegar la URL exacta]",
     "[completar]; cifras reportadas al 25-sep-2026",
     "Saldo del último mes del trimestre. Real = B × IPC(jun-2026) / IPC(mes). Índice = 100 × real / real(2021-I)."),
    ("Cartera vencida de consumo (V)", "Superintendencia Financiera de Colombia",
     "CONSUMO - VENCIDA (hoja C5)",
     "Igual a la cartera bruta (misma hoja y cobertura)",
     "Miles de millones de pesos corrientes",
     "Mensual; se usan mar/jun/sep/dic 2021-2026",
     "Cartera_hasta_agosto_Superfinanciera.xlsx, hoja C5",
     "https://www.superfinanciera.gov.co  [pegar la URL exacta]",
     "[completar]; cifras reportadas al 25-sep-2026",
     "Vencimiento: mora > 30 días (criterio SFC; confirmar en la Circular Básica Contable y Financiera, Cap. II). "
     "No se usa DETERIORO (provisiones) ni calidad por calificación (cartera riesgosa B-E). "
     "Real e índice igual que B. Morosidad = 100 × V / B."),
    ("Consumo final de los hogares (C)", "DANE",
     "Gasto de consumo final individual de los hogares; gasto de consumo final de las ISFLSH "
     "(P.311 + P.313 + P.323), Cuadro 2",
     "Total nacional; incluye ISFLSH; excluye Gobierno",
     "Miles de millones de pesos, series encadenadas de volumen, año de referencia 2015; "
     "ajustada por efecto estacional y calendario",
     "Trimestral, 2020-IV–2026-II (2024p provisional; 2025-2026pr preliminar)",
     "anex-GastoConstantes-IItrim2026.xlsx, Cuadro 2",
     "https://www.dane.gov.co  (Cuentas nacionales trimestrales, PIB enfoque del gasto)  [pegar la URL exacta]",
     "[completar]; anexo actualizado el 18-ago-2026",
     "No se deflacta. Índice = 100 × C / C(2021-I). Variación trimestral = 100 × (C_t / C_t-1 − 1)."),
    ("IPC (deflactor)", "DANE (vía Banco de la República)",
     "Índice de Precios al Consumidor (IPC) (Dato fin de mes)",
     "Total nacional", "Índice, base dic-2018 = 100",
     "Mensual, ene-2021–jun-2026", "Datos_BANREP_2.xlsx, hoja Datos",
     "https://www.banrep.gov.co/es/estadisticas  [pegar la URL exacta]",
     "16-oct-2026 (pie del archivo)",
     "IPC del último mes de cada trimestre. IPC base fijo = jun-2026 (159,53)."),
    ("Inflación anual", "DANE (vía Banco de la República)",
     "Índice de Precios al Consumidor (IPC) (Variación porcentual anual)",
     "Total nacional", "% (variación a 12 meses)",
     "Mensual, ene-2021–jun-2026", "Datos_BANREP_2.xlsx, hoja Datos",
     "https://www.banrep.gov.co/es/estadisticas  [pegar la URL exacta]",
     "16-oct-2026 (pie del archivo)",
     "Dato del último mes de cada trimestre (mar, jun, sep, dic). Idéntica a 'Inflación total anual'."),
]

ENC_BASE = [
    ("Trimestre", "", 11, None),
    ("TPM", "% e.a., prom. trim.", 10, "0.00"),
    ("Tasa crédito de consumo", "% e.a., prom. trim.", 12, "0.00"),
    ("Diferencial consumo − TPM", "p.p.", 11, "0.00"),
    ("Inflación anual", "%, último mes", 10, "0.00"),
    ("IPC", "dic-2018=100, último mes", 11, "0.00"),
    ("Cartera bruta consumo, nominal (B)", "miles de millones $ corrientes", 15, "#,##0.0"),
    ("Cartera vencida consumo, nominal (V)", "miles de millones $ corrientes", 15, "#,##0.0"),
    ("Cartera bruta real", "miles de millones $ de jun-2026", 14, "#,##0.0"),
    ("Cartera vencida real", "miles de millones $ de jun-2026", 14, "#,##0.0"),
    ("Índice cartera bruta real", "2021-I = 100", 11, "0.0"),
    ("Índice cartera vencida real", "2021-I = 100", 11, "0.0"),
    ("Morosidad (100×V/B)", "%", 11, "0.00"),
    ("Consumo final hogares (C)", "miles de millones $, volumen ref. 2015, desest.", 16, "#,##0.0"),
    ("Índice consumo hogares", "2021-I = 100", 11, "0.0"),
    ("Variación trimestral consumo", "%", 11, "0.00"),
]


def _filas_trim(fila_mes, trim):
    """Filas de la hoja Mensual del primer y último mes del trimestre (p. ej. '2021-I')."""
    filas = [f for mes, f in fila_mes.items() if etiqueta_trim(mes.asfreq("Q")) == trim]
    return filas[0], filas[-1]


def _hoja_tpm(wb, diaria):
    ws = wb.create_sheet("TPM_diaria")
    _fmt(ws, "A1", "Fecha", NEGRITA)
    _fmt(ws, "B1", "TPM diaria (%)", NEGRITA)
    d = diaria[diaria["fecha"] <= "2026-06-30"]
    for i, (f, v) in enumerate(zip(d["fecha"], d["tpm"]), start=2):
        _fmt(ws, f"A{i}", f.to_pydatetime(), AZUL_DATO, "dd/mm/yyyy")
        _fmt(ws, f"B{i}", v, AZUL_DATO, "0.00")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A2"
    return len(d) + 1


def _hoja_mensual(wb, m, ult_tpm):
    ws = wb.create_sheet("Mensual")
    enc = [("Mes", 10, "mmm-yyyy"), ("Trimestre", 10, None),
           ("TPM, promedio diario del mes (%)", 14, "0.000"), ("TPM 'Dato mes' BanRep (%) [control]", 14, "0.00"),
           ("Tasa crédito de consumo, prom. mensual (% e.a.)", 16, "0.00"), ("IPC (dic-2018=100)", 11, "0.00"),
           ("Inflación anual (%)", 11, "0.00"), ("Cartera bruta consumo (miles de millones $)", 16, "#,##0.0"),
           ("Cartera vencida consumo (miles de millones $)", 16, "#,##0.0"),
           ("Indicador calidad tradicional consumo SFC (fracción) [control]", 18, "0.000000")]
    for j, (t, w, _) in enumerate(enc, start=1):
        c = ws.cell(1, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 58
    rango_f, rango_v = f"TPM_diaria!$A$2:$A${ult_tpm}", f"TPM_diaria!$B$2:$B${ult_tpm}"
    for i, (mes, r) in enumerate(m.iterrows(), start=2):
        _fmt(ws, f"A{i}", mes.to_timestamp().to_pydatetime(), NEGRO, enc[0][2])
        _fmt(ws, f"B{i}", r["trimestre"], NEGRO)
        _fmt(ws, f"C{i}", f'=AVERAGEIFS({rango_v},{rango_f},">="&A{i},{rango_f},"<="&EOMONTH(A{i},0))',
             VERDE, enc[2][2])
        for col, campo in zip("DEFGHIJ", ["tpm_dato_mes", "tasa_consumo", "ipc", "inflacion_anual",
                                          "cartera_bruta", "cartera_vencida", "ind_calidad_sfc"]):
            _fmt(ws, f"{col}{i}", float(r[campo]), AZUL_DATO, enc["ABCDEFGHIJ".index(col)][2])
    ws.freeze_panes = "C2"
    return {mes: i for i, mes in enumerate(m.index, start=2)}


def _hoja_base(wb, q, fila_mes, ipc_base_mes):
    ws = wb.create_sheet("Base_trimestral", 1)
    _fmt(ws, "A1", "Base trimestral: crédito, consumo y morosidad en Colombia, 2021-I a 2026-II", TITULO)
    _fmt(ws, "A2", "2020-IV se incluye solo para calcular la primera variación del consumo. "
                   "Azul = dato transcrito de la fuente; negro = fórmula; verde = vínculo a otra hoja.", GRIS)
    _fmt(ws, "A4", "IPC base fijo (jun-2026)", NEGRITA)
    c = _fmt(ws, "C4", f"=Mensual!F{fila_mes[ipc_base_mes]}", VERDE, "0.00")
    c.fill = RELLENO_PARAM
    c.comment = Comment("Parámetro: IPC de junio de 2026 (DANE). Los saldos reales quedan en pesos de jun-2026. "
                        "Cambiar el mes base no altera los índices 2021-I = 100.", "Taller 2")
    f_enc, f_uni, f0 = 6, 7, 8
    for j, (t, u, w, _) in enumerate(ENC_BASE, start=1):
        for fila, txt, fuente in ((f_enc, t, NEGRITA), (f_uni, u, GRIS)):
            c = ws.cell(fila, j, txt)
            c.font, c.alignment = fuente, CENTRO
            c.fill = RELLENO_ENC
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[f_enc].height = 44
    ws.row_dimensions[f_uni].height = 44
    fila_base = f0 + list(q.index).index(TRIM_BASE)
    for k, (trim, r) in enumerate(q.iterrows()):
        i = f0 + k
        _fmt(ws, f"A{i}", trim, NEGRITA)
        fmt = {get_column_letter(j): e[3] for j, e in enumerate(ENC_BASE, start=1)}
        _fmt(ws, f"N{i}", float(r["consumo_hogares"]), AZUL_DATO, fmt["N"])
        if k > 0:
            _fmt(ws, f"P{i}", f"=100*(N{i}/N{i - 1}-1)", NEGRO, fmt["P"])
            _fmt(ws, f"O{i}", f"=100*N{i}/N${fila_base}", NEGRO, fmt["O"])
        else:
            _fmt(ws, f"O{i}", f"=100*N{i}/N${fila_base}", NEGRO, fmt["O"])
        if trim == "2020-IV":
            continue
        r1, r3 = _filas_trim(fila_mes, trim)
        _fmt(ws, f"B{i}", f"=AVERAGE(Mensual!C{r1}:C{r3})", VERDE, fmt["B"])
        _fmt(ws, f"C{i}", f"=AVERAGE(Mensual!E{r1}:E{r3})", VERDE, fmt["C"])
        _fmt(ws, f"D{i}", f"=C{i}-B{i}", NEGRO, fmt["D"])
        _fmt(ws, f"E{i}", f"=Mensual!G{r3}", VERDE, fmt["E"])
        _fmt(ws, f"F{i}", f"=Mensual!F{r3}", VERDE, fmt["F"])
        _fmt(ws, f"G{i}", f"=Mensual!H{r3}", VERDE, fmt["G"])
        _fmt(ws, f"H{i}", f"=Mensual!I{r3}", VERDE, fmt["H"])
        _fmt(ws, f"I{i}", f"=G{i}*$C$4/F{i}", NEGRO, fmt["I"])
        _fmt(ws, f"J{i}", f"=H{i}*$C$4/F{i}", NEGRO, fmt["J"])
        _fmt(ws, f"K{i}", f"=100*I{i}/I${fila_base}", NEGRO, fmt["K"])
        _fmt(ws, f"L{i}", f"=100*J{i}/J${fila_base}", NEGRO, fmt["L"])
        _fmt(ws, f"M{i}", f"=100*H{i}/G{i}", NEGRO, fmt["M"])
    ult = f0 + len(q) - 1
    for j in range(1, len(ENC_BASE) + 1):
        ws.cell(ult, j).border = BORDE
    notas = [
        "Notas metodológicas:",
        "• Tasas: TPM = promedio de la serie diaria por mes y luego de los 3 meses; tasa de consumo = promedio de las "
        "3 tasas mensuales (créditos nuevos, no incluye tarjetas). Ambas en % efectivo anual.",
        "• Cartera: saldo del último mes de cada trimestre (SFC, hoja C5), establecimientos de crédito sin IOE, "
        "incluye leasing. No se suman saldos mensuales. Vencida ≠ riesgosa (calificación B-E) ≠ provisiones (deterioro).",
        "• Real = nominal × IPC(jun-2026) / IPC(mes del saldo). Morosidad = 100 × V / B (con saldos nominales; "
        "es igual con reales porque el deflactor se cancela).",
        "• Consumo: gasto de consumo final de los hogares incluidas ISFLSH, volumen, desestacionalizado; no se "
        "deflacta de nuevo. 2020-IV solo alimenta la variación de 2021-I.",
        "• Limitación de cobertura: la cartera SFC incluye tarjetas de crédito; la tasa de consumo BanRep no. "
        "Para crédito y morosidad se prioriza la SFC.",
        "• Sin datos faltantes en la ventana 2021-I a 2026-II: no hubo que dejar celdas vacías ni interpolar.",
    ]
    for k, t in enumerate(notas):
        _fmt(ws, f"A{ult + 2 + k}", t, NEGRITA if k == 0 else GRIS)
    ws.freeze_panes = "B8"
    return f0, ult


def _hoja_ficha(wb):
    ws = wb.create_sheet("Ficha_fuentes", 2)
    _fmt(ws, "A1", "Ficha de fuentes", TITULO)
    _fmt(ws, "A2", "Los campos entre corchetes [ ] deben completarse con la URL exacta y la fecha en que se "
                   "consultó cada fuente; no se inventan.", GRIS)
    enc = ["Variable", "Entidad", "Nombre exacto de la serie", "Cobertura", "Unidad", "Frecuencia y período",
           "Archivo original (data/raw)", "Enlace", "Fecha de consulta", "Transformaciones"]
    anchos = [20, 18, 34, 28, 26, 22, 26, 34, 20, 46]
    for j, (t, w) in enumerate(zip(enc, anchos), start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, fila in enumerate(FICHA, start=5):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(i, j, v)
            c.font, c.alignment = NEGRO, AJUSTE
        ws.row_dimensions[i].height = 105
    ws.freeze_panes = "B5"


def _hoja_controles(wb, q, f0, fila_mes):
    ws = wb.create_sheet("Controles")
    _fmt(ws, "A1", "Controles de cálculo (diferencias ≈ 0 confirman que las fórmulas reproducen las fuentes)", TITULO)
    enc = ["Trimestre", "Morosidad propia (%)", "Indicador SFC × 100 (%)", "Diferencia",
           "Var. trim. consumo propia (%)", "Var. trim. publicada DANE (%)", "Diferencia",
           "TPM prom. diario (%)", "TPM 'Dato mes' prom. (%)", "Diferencia (p.p.)"]
    for j, t in enumerate(enc, start=1):
        c = ws.cell(3, j, t)
        c.font, c.fill, c.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.row_dimensions[3].height = 44
    i = 4
    for k, (trim, r) in enumerate(q.iterrows()):
        if trim == "2020-IV":
            continue
        fb = f0 + k
        r1, r3 = _filas_trim(fila_mes, trim)
        _fmt(ws, f"A{i}", trim, NEGRITA)
        _fmt(ws, f"B{i}", f"=Base_trimestral!M{fb}", VERDE, "0.0000")
        _fmt(ws, f"C{i}", f"=100*Mensual!J{r3}", VERDE, "0.0000")
        _fmt(ws, f"D{i}", f"=B{i}-C{i}", NEGRO, "0.0000")
        _fmt(ws, f"E{i}", f"=Base_trimestral!P{fb}", VERDE, "0.0000")
        _fmt(ws, f"F{i}", float(r["var_dane"]), AZUL_DATO, "0.0000")
        _fmt(ws, f"G{i}", f"=E{i}-F{i}", NEGRO, "0.0000")
        _fmt(ws, f"H{i}", f"=Base_trimestral!B{fb}", VERDE, "0.000")
        _fmt(ws, f"I{i}", f"=AVERAGE(Mensual!D{r1}:D{r3})", VERDE, "0.000")
        _fmt(ws, f"J{i}", f"=H{i}-I{i}", NEGRO, "0.000")
        i += 1
    _fmt(ws, f"A{i + 1}", "Máx. |dif.|", NEGRITA)
    for col in "DGJ":
        _fmt(ws, f"{col}{i + 1}", f"=MAX(MAX({col}4:{col}{i - 1}),-MIN({col}4:{col}{i - 1}))", NEGRITA, "0.0000")
    _fmt(ws, f"A{i + 3}", "D: el indicador SFC viene redondeado a 6 decimales (fracción), por eso la diferencia "
                         "es del orden de 0,0001 p.p. G: la variación publicada por el DANE en el Cuadro 2. "
                         "J: diferencia por la forma de agregar la TPM (decisión 1); máximo 0,15 p.p. en 2022-IV.",
         GRIS)


def _hoja_figuras(wb, rutas):
    ws = wb.create_sheet("Figuras", 3)
    _fmt(ws, "A1", "Figuras del taller (también en output/figuras/*.png)", TITULO)
    fila = 3
    for ruta in rutas:
        img = XLImage(str(ruta))
        img.width, img.height = img.width * 0.36, img.height * 0.36
        ws.add_image(img, f"A{fila}")
        fila += int(img.height / 20) + 3


def _hoja_leeme(wb, ipc_base):
    ws = wb.active
    ws.title = "Leeme"
    ws.column_dimensions["A"].width = 120
    lineas = [
        ("Taller 2 – Política Económica (Univalle): crédito, consumo y morosidad en Colombia", TITULO),
        ("Base trimestral 2021-I a 2026-II construida con scripts/construir_base.py a partir de data/raw/.", NEGRO),
        ("", NEGRO),
        ("Hojas", NEGRITA),
        ("• Base_trimestral: la base del taller (todas las series, incluida la inflación). Fórmulas vivas.", NEGRO),
        ("• Ficha_fuentes: entidad, nombre exacto, cobertura, unidad, enlace, fecha de consulta y transformaciones.",
         NEGRO),
        ("• Figuras: las 4 figuras pedidas (título, ejes, unidad, período, fuente y nota).", NEGRO),
        ("• Mensual: insumos mensuales transcritos de las fuentes; la TPM mensual se calcula desde TPM_diaria.", NEGRO),
        ("• TPM_diaria: serie diaria de la TPM (BanRep) que alimenta el promedio mensual.", NEGRO),
        ("• Controles: verifica que los cálculos reproducen los indicadores publicados (SFC, DANE).", NEGRO),
        ("", NEGRO),
        ("Convención de colores", NEGRITA),
        ("Azul = dato transcrito de la fuente · Negro = fórmula en la misma hoja · Verde = vínculo a otra hoja · "
         "Amarillo = parámetro (IPC base).", NEGRO),
        ("", NEGRO),
        ("Decisiones metodológicas (aprobadas por el grupo)", NEGRITA),
        ("1. TPM mensual = promedio de la serie diaria del mes; trimestral = promedio de los 3 meses.", NEGRO),
        ("2. Tasa de consumo BanRep excluye tarjetas; cartera SFC las incluye → limitación documentada; "
         "se prioriza la SFC para crédito y morosidad.", NEGRO),
        ("3. Consumo de los hogares: DANE, Cuadro 2 (volumen, desestacionalizado), incluye ISFLSH.", NEGRO),
        (f"4. IPC base fijo = jun-2026 ({ipc_base:.2f}): saldos reales en pesos de junio de 2026.", NEGRO),
        ("", NEGRO),
        ("Datos faltantes: ninguno en la ventana. No se interpoló ni se imputó ninguna observación.", NEGRO),
        ("Pendientes antes de entregar: completar URL exacta y fecha de consulta en Ficha_fuentes; citar la "
         "definición de vencimiento de la SFC (Circular Básica Contable y Financiera, Cap. II).", NEGRO),
    ]
    for i, (t, f) in enumerate(lineas, start=1):
        _fmt(ws, f"A{i}", t, f)


def exportar(diaria, m, q, ipc_base):
    PROC.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    m_csv = m.copy()
    m_csv.index = m_csv.index.astype(str)
    m_csv.to_csv(PROC / "insumos_mensuales.csv", index_label="mes", float_format="%.6f")
    q.drop(columns=["var_dane"]).to_csv(PROC / "base_trimestral.csv", index_label="trimestre",
                                        float_format="%.6f")
    rutas = figuras(q)

    wb = Workbook()
    _hoja_leeme(wb, ipc_base)
    ult_tpm = _hoja_tpm(wb, diaria)
    fila_mes = _hoja_mensual(wb, m, ult_tpm)
    f0, _ = _hoja_base(wb, q, fila_mes, MES_BASE_IPC)
    _hoja_ficha(wb)
    _hoja_figuras(wb, rutas)
    _hoja_controles(wb, q, f0, fila_mes)
    wb.save(OUT / "Taller2_base_trimestral.xlsx")
    print("Listo:", OUT / "Taller2_base_trimestral.xlsx")
