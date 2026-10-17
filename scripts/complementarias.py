"""Figuras y tabla complementarias (extensiones opcionales del taller).

E1. Tasas reales ex post (TPM y consumo) y diferencial consumo − TPM.
E2. Cartera bruta real vs. consumo real de los hogares (niveles y crecimiento anual).
E3. Descomposición de la variación anual de la morosidad: vencida vs. bruta.
E4. Consumo de los hogares por durabilidad (DANE, Cuadro 4).

Todo sale de data/raw/ y de la base ya construida; no se agregan fuentes nuevas.
"""
from __future__ import annotations

import numpy as np
import openpyxl
import pandas as pd

from construir_base import ARCH_DANE, TRIM_BASE, etiqueta_trim

DURABILIDAD = ["Bienes durables", "Bienes semidurables", "Bienes no durables", "Servicios"]


def leer_durabilidad() -> pd.DataFrame:
    """Cuadro 4 (volumen, desestacionalizado): primera tabla = niveles por durabilidad."""
    filas = list(openpyxl.load_workbook(ARCH_DANE, data_only=True)["Cuadro 4"].iter_rows(values_only=True))
    i_anio = next(i for i, f in enumerate(filas) if f[0] == "Divisiones COICOP1")
    fin = next(i for i, f in enumerate(filas) if f[0] and str(f[0]).startswith("Fuente"))
    anios, trims = filas[i_anio], filas[i_anio + 1]
    cols, anio = {}, None
    for j in range(2, len(trims)):
        if anios[j]:
            anio = int(str(anios[j])[:4])
        if trims[j]:
            q = {"I": 1, "II": 2, "III": 3, "IV": 4}[trims[j]]
            cols[etiqueta_trim(pd.Period(year=anio, quarter=q, freq="Q"))] = j
    df = pd.DataFrame(index=list(cols))
    for f in filas[i_anio:fin]:
        if f[1] in DURABILIDAD:
            df[f[1]] = [float(f[j]) for j in cols.values()]
    return df[DURABILIDAD]


def calcular(q: pd.DataFrame) -> pd.DataFrame:
    """Series complementarias por trimestre (2021-I a 2026-II)."""
    q = q.loc[TRIM_BASE:]  # 2020-IV queda fuera: solo sirve para la primera variación trimestral
    c = pd.DataFrame(index=q.index)
    pi = q["inflacion_anual"] / 100
    # Tasa real ex post (Fisher exacta): (1 + i) / (1 + π) − 1, con π = inflación anual del último mes
    c["tpm_real"] = 100 * ((1 + q["tpm"] / 100) / (1 + pi) - 1)
    c["consumo_real_tasa"] = 100 * ((1 + q["tasa_consumo"] / 100) / (1 + pi) - 1)
    c["diferencial"] = q["diferencial"]
    c["inflacion_anual"] = q["inflacion_anual"]
    # Crecimiento anual (t frente a t−4); el primer dato posible es 2022-I
    c["var_anual_bruta_real"] = 100 * (q["bruta_real"] / q["bruta_real"].shift(4) - 1)
    c["var_anual_consumo"] = 100 * (q["consumo_hogares"] / q["consumo_hogares"].shift(4) - 1)
    # Descomposición exacta en logaritmos: Δln(M) = Δln(V) − Δln(B)
    c["dln_vencida"] = 100 * np.log(q["vencida_nom"] / q["vencida_nom"].shift(4))
    c["menos_dln_bruta"] = -100 * np.log(q["bruta_nom"] / q["bruta_nom"].shift(4))
    c["dln_morosidad"] = 100 * np.log(q["morosidad"] / q["morosidad"].shift(4))
    c["idx_bruta_real"] = q["idx_bruta_real"]
    c["idx_consumo"] = q["idx_consumo"]
    dur = leer_durabilidad()
    for k in DURABILIDAD:
        c[k] = dur.loc[c.index, k]
        c[f"idx_{k}"] = 100 * c[k] / c.loc[TRIM_BASE, k]
    return c


# ----------------------------------------------------------------------------- figuras E1–E4
def figuras_extra(c: pd.DataFrame) -> list:
    import matplotlib.pyplot as plt

    from exportar import FIG, REJILLA, SERIE1, SERIE2, TINTA, TINTA2, _coma, _ejes, _etiqueta_final, _marco

    serie3, serie4 = "#1baf7a", "#eda100"  # 3.º y 4.º color de la paleta categórica validada
    et = list(c.index)
    x = np.arange(len(et))
    rutas = []

    def etiquetas_finales(ax, items, separacion):
        """Etiquetas al final de cada línea, separadas verticalmente para que no se monten."""
        items = sorted(items, key=lambda t: t[0])
        pos = [items[0][0]]
        for y, *_ in items[1:]:
            pos.append(max(y, pos[-1] + separacion))
        for (y, texto, color), yp in zip(items, pos):
            ax.plot([x[-1]], [y], "o", color=color, markersize=5, markeredgecolor="white", markeredgewidth=1.5)
            ax.annotate(texto, (x[-1], yp), xytext=(8, 0), textcoords="offset points", va="center",
                        fontsize=9, color=TINTA, fontweight="bold", annotation_clip=False)

    def guardar(fig, nombre):
        rutas.append(FIG / nombre)
        fig.savefig(rutas[-1], dpi=200)
        plt.close(fig)

    # E1 ---------------------------------------------------------------------
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(10, 8), sharex=True, gridspec_kw={"height_ratios": [3, 2]})
    fig.subplots_adjust(top=0.865, bottom=0.15, left=0.08, right=0.92, hspace=0.25)
    a1.axhline(0, color="#8a8984", lw=1)
    a1.plot(x, c["consumo_real_tasa"], color=SERIE2, lw=2, label="Tasa real del crédito de consumo")
    a1.plot(x, c["tpm_real"], color=SERIE1, lw=2, label="Tasa real de intervención (TPM)")
    a1.plot(x, c["inflacion_anual"], color=serie3, lw=1.6, ls="--", label="Inflación anual (referencia)")
    for col, color in (("consumo_real_tasa", SERIE2), ("tpm_real", SERIE1)):
        _etiqueta_final(a1, x[-1], c[col].iloc[-1], _coma(c[col].iloc[-1], 1), color)
    neg = np.where(c["tpm_real"].values <= 0)[0]
    i0 = int(neg[-1]) + 1  # primer trimestre de la racha positiva actual
    a1.annotate(f"TPM real positiva desde {et[i0]}", (i0, c["tpm_real"].iloc[i0]), xytext=(25, -30),
                textcoords="offset points", ha="left", fontsize=8.5, color=TINTA2,
                arrowprops=dict(arrowstyle="-", color=TINTA2, lw=0.8))
    a1.set_title("A. Tasas reales ex post (%)", loc="left", fontsize=10, color=TINTA)
    a1.set_ylabel("%", color=TINTA2)
    a1.legend(loc="upper left", frameon=False, fontsize=8.5, ncol=3)
    a1.set_ylim(-7, 21)
    a2.bar(x, c["diferencial"], color=SERIE2, width=0.7)
    i_max = int(c["diferencial"].values.argmax())
    for i in sorted({0, i_max, len(et) - 1}):
        a2.annotate(_coma(c["diferencial"].iloc[i], 1), (i, c["diferencial"].iloc[i]), xytext=(0, 3),
                    textcoords="offset points", ha="center", fontsize=8.5, color=TINTA, fontweight="bold")
    a2.set_title("B. Diferencial: tasa de consumo − TPM (puntos porcentuales)", loc="left", fontsize=10, color=TINTA)
    a2.set_ylabel("p.p.", color=TINTA2)
    a2.set_ylim(0, 20)
    _ejes(a1, et)
    _ejes(a2, et)
    _marco(fig, "Figura E1. Costo real del crédito y traslado de la tasa de política",
           "Promedios trimestrales de tasas nominales deflactadas con la inflación anual, 2021-I a 2026-II",
           "Fuente: Banco de la República (TPM, tasa de consumo sin tarjetas); DANE (IPC). Cálculos propios.",
           "Tasa real ex post = 100 × [(1 + i) / (1 + π) − 1], con i en % e.a. y π = inflación anual del último mes "
           "del trimestre. Mide el costo real ya observado, no el esperado. El diferencial es nominal.")
    guardar(fig, "figuraE1_tasas_reales.png")

    # E2 ---------------------------------------------------------------------
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(10, 8), sharex=True)
    fig.subplots_adjust(top=0.865, bottom=0.15, left=0.08, right=0.92, hspace=0.25)
    a1.axhline(100, color="#8a8984", lw=1, ls="--")
    a1.plot(x, c["idx_consumo"], color=SERIE2, lw=2, label="Consumo real de los hogares")
    a1.plot(x, c["idx_bruta_real"], color=SERIE1, lw=2, label="Cartera bruta de consumo, real")
    for col, color in (("idx_consumo", SERIE2), ("idx_bruta_real", SERIE1)):
        _etiqueta_final(a1, x[-1], c[col].iloc[-1], _coma(c[col].iloc[-1], 1), color)
    a1.set_title("A. Niveles, índice 2021-I = 100", loc="left", fontsize=10, color=TINTA)
    a1.set_ylabel("Índice", color=TINTA2)
    a1.legend(loc="upper left", frameon=False, fontsize=9)
    a2.axhline(0, color="#8a8984", lw=1)
    a2.plot(x, c["var_anual_consumo"], color=SERIE2, lw=2, marker="o", markersize=4)
    a2.plot(x, c["var_anual_bruta_real"], color=SERIE1, lw=2, marker="o", markersize=4)
    i_min = int(np.nanargmin(c["var_anual_bruta_real"].values))
    a2.annotate(_coma(c["var_anual_bruta_real"].iloc[i_min], 1), (i_min, c["var_anual_bruta_real"].iloc[i_min]),
                xytext=(0, -14), textcoords="offset points", ha="center", fontsize=8.5, color=TINTA,
                fontweight="bold")
    etiquetas_finales(a2, [(c[col].iloc[-1], _coma(c[col].iloc[-1], 1), color)
                           for col, color in (("var_anual_consumo", SERIE2), ("var_anual_bruta_real", SERIE1))], 2.2)
    a2.set_ylim(-15, 19)
    a2.set_title("B. Crecimiento anual (%, frente al mismo trimestre del año anterior; desde 2022-I)",
                 loc="left", fontsize=10, color=TINTA)
    a2.set_ylabel("%", color=TINTA2)
    _ejes(a1, et)
    _ejes(a2, et)
    _marco(fig, "Figura E2. Crédito de consumo y consumo de los hogares: ¿se mueven juntos?",
           "Cartera bruta real (SFC, deflactada con IPC) y consumo final de los hogares en volumen (DANE), 2021-I a 2026-II",
           "Fuente: Superintendencia Financiera (cartera); DANE (IPC y PIB por el gasto). Cálculos propios.",
           "Cartera: saldo del último mes del trimestre en pesos de jun-2026. Consumo: incluye ISFLSH, desestacionalizado. "
           "El crecimiento anual empieza en 2022-I porque la base empieza en 2021-I (no se usan datos de 2020).")
    guardar(fig, "figuraE2_credito_vs_consumo.png")

    # E3 ---------------------------------------------------------------------
    d = c.dropna(subset=["dln_morosidad"])
    et3, x3 = list(d.index), np.arange(len(d))
    fig, ax = plt.subplots(figsize=(10, 6.2))
    fig.subplots_adjust(top=0.84, bottom=0.22, left=0.08, right=0.95)
    v, b = d["dln_vencida"].values, d["menos_dln_bruta"].values
    base_pos = np.zeros(len(d))
    base_neg = np.zeros(len(d))
    for vals, color, nombre in ((v, SERIE2, "Aporte de la cartera vencida (numerador)"),
                                (b, SERIE1, "Aporte de la cartera bruta (denominador, con signo −)")):
        fondo = np.where(vals >= 0, base_pos, base_neg)
        ax.bar(x3, vals, bottom=fondo, color=color, width=0.7, label=nombre, edgecolor="white", linewidth=1)
        base_pos += np.where(vals >= 0, vals, 0)
        base_neg += np.where(vals < 0, vals, 0)
    ax.plot(x3, d["dln_morosidad"], color=TINTA, lw=1.6, marker="o", markersize=5, markeredgecolor="white",
            label="Variación de la morosidad")
    ax.axhline(0, color="#8a8984", lw=1)
    ax.set_ylabel("Puntos logarítmicos × 100 (≈ % de cambio anual)", color=TINTA2)
    ax.legend(loc="upper right", frameon=False, fontsize=8.5)
    _ejes(ax, et3)
    _marco(fig, "Figura E3. ¿Por qué cambia la morosidad? Numerador vs. denominador",
           "Descomposición de la variación anual del índice de morosidad de consumo, 2022-I a 2026-II",
           "Fuente: Superintendencia Financiera de Colombia. Cálculos propios.",
           "Δln(M) = Δln(V) − Δln(B), comparando cada trimestre con el mismo del año anterior; la suma de las barras es "
           "exactamente la línea. Barra naranja > 0: sube el saldo vencido; barra azul < 0: crece la cartera bruta "
           "(baja la morosidad); barra azul > 0: se contrae la cartera bruta (sube la morosidad).")
    guardar(fig, "figuraE3_descomposicion_morosidad.png")

    # E4 ---------------------------------------------------------------------
    fig, ax = plt.subplots(figsize=(10, 6.2))
    fig.subplots_adjust(top=0.86, bottom=0.2, left=0.07, right=0.84)
    ax.axhline(100, color="#8a8984", lw=1, ls="--")
    cortos = {"Bienes durables": "durables", "Bienes semidurables": "semidurables",
              "Bienes no durables": "no durables", "Servicios": "servicios"}
    items = []
    for k, color in zip(DURABILIDAD, (SERIE1, SERIE2, serie3, serie4)):
        ax.plot(x, c[f"idx_{k}"], color=color, lw=2, label=k)
        items.append((c[f"idx_{k}"].iloc[-1], f"{_coma(c[f'idx_{k}'].iloc[-1], 1)}  {cortos[k]}", color))
    etiquetas_finales(ax, items, 2.6)
    ax.set_ylabel("Índice, 2021-I = 100", color=TINTA2)
    ax.legend(loc="upper left", frameon=False, fontsize=9)
    ax.grid(axis="y", color=REJILLA)
    _ejes(ax, et)
    _marco(fig, "Figura E4. Consumo de los hogares por durabilidad",
           "Volumen (series encadenadas, referencia 2015), ajustado por estacionalidad y calendario; índice 2021-I = 100",
           "Fuente: DANE, PIB trimestral por el enfoque del gasto, Cuadro 4 (actualizado 18-ago-2026). Cálculos propios.",
           "Gasto de los hogares en el territorio económico (concepto distinto al total con ISFLSH de la Figura 4). "
           "Las series encadenadas no son aditivas: no se deben sumar los componentes.")
    guardar(fig, "figuraE4_consumo_durabilidad.png")
    return rutas


# ----------------------------------------------------------------------------- hoja Excel
def hoja_complementarias(wb, c: pd.DataFrame, f0_base: int, rutas: list):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    from exportar import AZUL_DATO, CENTRO, GRIS, NEGRITA, NEGRO, RELLENO_ENC, TITULO, VERDE, _fmt

    ws = wb.create_sheet("Complementarias", 4)
    _fmt(ws, "A1", "Series y figuras complementarias (extensiones opcionales; no reemplazan las 4 figuras)", TITULO)
    _fmt(ws, "A2", "Fórmulas sobre Base_trimestral. Crecimientos anuales desde 2022-I (no se usan datos de 2020). "
                   "Azul = dato DANE Cuadro 4; verde = vínculo; negro = fórmula.", GRIS)
    enc = [("Trimestre", None), ("TPM real ex post (%)", "0.00"), ("Tasa de consumo real ex post (%)", "0.00"),
           ("Diferencial consumo − TPM (p.p.)", "0.00"), ("Var. anual cartera bruta real (%)", "0.00"),
           ("Var. anual consumo hogares (%)", "0.00"), ("Aporte vencida: 100·Δln V", "0.00"),
           ("Aporte bruta: −100·Δln B", "0.00"), ("Var. morosidad: 100·Δln M", "0.00"),
           ("Control (G+H−I)", "0.0000")]
    enc += [(f"{k} (miles de millones $, vol.)", "#,##0.0") for k in DURABILIDAD]
    enc += [(f"Índice {k.lower()} (2021-I=100)", "0.0") for k in DURABILIDAD]
    for j, (t, _) in enumerate(enc, start=1):
        cel = ws.cell(4, j, t)
        cel.font, cel.fill, cel.alignment = NEGRITA, RELLENO_ENC, CENTRO
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[4].height = 58
    fb0 = f0_base + 1  # fila de 2021-I en Base_trimestral (la primera es 2020-IV)
    r0 = 5
    for k, trim in enumerate(c.index):
        r, b = r0 + k, fb0 + k
        B = "Base_trimestral!"
        _fmt(ws, f"A{r}", f"={B}A{b}", VERDE)
        _fmt(ws, f"B{r}", f"=100*((1+{B}B{b}/100)/(1+{B}E{b}/100)-1)", VERDE, "0.00")
        _fmt(ws, f"C{r}", f"=100*((1+{B}C{b}/100)/(1+{B}E{b}/100)-1)", VERDE, "0.00")
        _fmt(ws, f"D{r}", f"={B}D{b}", VERDE, "0.00")
        if k >= 4:
            _fmt(ws, f"E{r}", f"=100*({B}I{b}/{B}I{b - 4}-1)", VERDE, "0.00")
            _fmt(ws, f"F{r}", f"=100*({B}N{b}/{B}N{b - 4}-1)", VERDE, "0.00")
            _fmt(ws, f"G{r}", f"=100*LN({B}H{b}/{B}H{b - 4})", VERDE, "0.00")
            _fmt(ws, f"H{r}", f"=-100*LN({B}G{b}/{B}G{b - 4})", VERDE, "0.00")
            _fmt(ws, f"I{r}", f"=100*LN({B}M{b}/{B}M{b - 4})", VERDE, "0.00")
            _fmt(ws, f"J{r}", f"=G{r}+H{r}-I{r}", NEGRO, "0.0000")
        for j, kname in enumerate(DURABILIDAD):
            col_v, col_i = get_column_letter(11 + j), get_column_letter(15 + j)
            _fmt(ws, f"{col_v}{r}", float(c.loc[trim, kname]), AZUL_DATO, "#,##0.0")
            _fmt(ws, f"{col_i}{r}", f"=100*{col_v}{r}/{col_v}${r0}", NEGRO, "0.0")
    ult = r0 + len(c) - 1
    ws.freeze_panes = "B5"

    # Puntos de giro (fórmulas INDEX/MATCH sobre las bases)
    pg = ult + 3
    _fmt(ws, f"A{pg}", "Puntos de giro (para fechar etapas y rezagos)", NEGRITA)
    for j, t in enumerate(["Serie", "Criterio", "Valor", "Trimestre"], start=1):
        cel = ws.cell(pg + 1, j, t)
        cel.font, cel.fill, cel.alignment = NEGRITA, RELLENO_ENC, CENTRO
    fb_ini, fb_fin = fb0, fb0 + len(c) - 1
    rango = lambda col: f"Base_trimestral!${col}${fb_ini}:${col}${fb_fin}"
    trims_base = f"Base_trimestral!$A${fb_ini}:$A${fb_fin}"
    trims_comp = f"$A${r0}:$A${ult}"
    giros = [
        ("TPM (%)", "MAX", rango("B"), trims_base),
        ("Tasa de crédito de consumo (%)", "MAX", rango("C"), trims_base),
        ("Tasa de consumo real ex post (%)", "MAX", f"$C${r0}:$C${ult}", trims_comp),
        ("Índice cartera bruta real", "MAX", rango("K"), trims_base),
        ("Índice cartera bruta real", "MIN", rango("K"), trims_base),
        ("Índice cartera vencida real", "MAX", rango("L"), trims_base),
        ("Morosidad (%)", "MAX", rango("M"), trims_base),
        ("Morosidad (%)", "MIN", rango("M"), trims_base),
        ("Var. trimestral consumo hogares (%)", "MIN", rango("P"), trims_base),
        ("Var. anual cartera bruta real (%)", "MIN", f"$E${r0}:$E${ult}", trims_comp),
    ]
    for k, (serie, crit, rg, tr) in enumerate(giros):
        r = pg + 2 + k
        _fmt(ws, f"A{r}", serie, NEGRO)
        _fmt(ws, f"B{r}", "Máximo" if crit == "MAX" else "Mínimo", NEGRO)
        _fmt(ws, f"C{r}", f"={crit}({rg})", NEGRO, "0.00")
        _fmt(ws, f"D{r}", f"=INDEX({tr},MATCH(C{r},{rg},0))", NEGRO)
    ws.column_dimensions["A"].width = 30

    fila = pg + 2 + len(giros) + 2
    _fmt(ws, f"A{fila}", "Figuras complementarias (también en output/figuras/figuraE*.png)", NEGRITA)
    fila += 1
    for ruta in rutas:
        img = XLImage(str(ruta))
        img.width, img.height = img.width * 0.36, img.height * 0.36
        ws.add_image(img, f"A{fila}")
        fila += int(img.height / 20) + 3
